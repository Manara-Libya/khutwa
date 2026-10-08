"""Export the HTML deck to PowerPoint, with its click steps as PowerPoint animations.

    cd pitch && python3 build.py
    uv run --with playwright --with python-pptx python tools/export_pptx.py      # writes khutwa-pitch.pptx

Each slide is rendered in Chrome at 1920x1080, so fonts, doodles and Arabic shaping look exactly as in the deck.
The slide without its animated parts becomes the background; every element that appears on a click (data-step)
becomes its own transparent layer, placed where it sits on the slide, with the matching PowerPoint effect:

    .step (rise)        Float In        data-anim="pop"     Zoom          data-anim="fade"   Fade
    data-anim="draw"    Wipe            data-anim="type"    Wipe          data-anim="type-erase"  Wipe in, then out
    doodles that draw with their slide  Wipe, automatically when the slide opens
    data-add-class      the changed element fades in over its old look (the survey dots)

Videos stay real videos: a chaptered video (slide 9) is cut into one clip per chapter, and each click shows and
plays the next clip, which then holds on its last frame. Script lines go into the speaker notes. A click that only
moves the script (no visual change) gets an invisible step, so the clicks match the HTML deck one to one.
Text in the PowerPoint is a picture, not editable: edit src/index.tpl.html and export again.
"""
import asyncio, json, os, shutil, subprocess, sys, tempfile
from pathlib import Path

PITCH = Path(__file__).resolve().parent.parent
OUT = PITCH / "khutwa-pitch.pptx"
PX = 6350  # EMU per CSS pixel on a 1920x1080 slide (13.333 in)
CHROME = os.environ.get("CHROME", "/usr/bin/google-chrome")

# --------------------------------------------------------------------------- capture (in the page)
PAGE_JS = r"""
(() => {
  const sl = document.querySelector('.slide.active');
  const stepOf = (e) => +e.dataset.step || 0;
  const visual = [...sl.querySelectorAll('[data-step]')].filter((e) => !e.closest('aside'));
  const parentStep = (e) => { let p = e.parentElement; while (p && p !== sl) { if (p.hasAttribute('data-step')) return p; p = p.parentElement; } return null; };
  const isSvg = (e) => e.tagName.toLowerCase() === 'svg';
  // final look of every step: text typed out, counters at their value, doodles drawn
  sl.querySelectorAll('.typed').forEach((t) => { t.dataset.run = ''; if (t.dataset.text) t.textContent = t.dataset.text; });
  sl.querySelectorAll('[data-count]').forEach((e) => { e.textContent = e.dataset.count; });
  sl.querySelectorAll('svg.doodle.draw, svg.ill.trace').forEach((s) => s.classList.add('drawn'));
  sl.querySelectorAll('.video-wrap').forEach((w) => w.classList.add('has-video'));
  sl.querySelectorAll('video').forEach((v) => v.pause());

  let n = 0;
  const tag = (e) => { if (!e.dataset.capId) e.dataset.capId = String(++n); return e.dataset.capId; };
  const layers = [];
  for (const e of visual) {
    if (!e.classList.contains('step')) continue;          // a trigger only (a video that plays on that step)
    const k = stepOf(e), p = parentStep(e);
    let delay = +(e.dataset.delay || 0);
    if (p && stepOf(p) === k && !isSvg(e)) continue;     // part of its parent's picture
    const anim = e.dataset.anim || (isSvg(e) ? 'draw' : 'rise');
    if (anim === 'type' || anim === 'type-erase') {
      const t = e.querySelector('.typed') || e;
      layers.push({ id: tag(t), step: k, anim, delay, chars: (t.textContent || '').length, typedBox: tag(e) });
      continue;
    }
    layers.push({ id: tag(e), step: k, anim: p && stepOf(p) === k ? 'draw' : anim, delay });
  }
  // doodles that draw by themselves when the slide opens
  sl.querySelectorAll('svg.doodle.draw, svg.ill.trace').forEach((s) => {
    if (!s.closest('[data-step]')) layers.push({ id: tag(s), step: 0, anim: 'draw', delay: +(s.dataset.delay || 0), auto: true });
  });
  layers.forEach((l) => { if (l.step === 0) l.auto = true; });
  // a step that changes another element's look (data-add-class): that element again, in its new look
  const changes = [];
  for (const e of visual) {
    if (!e.dataset.addClass) continue;
    const t = document.querySelector(e.dataset.target);
    if (!t || t === e || e.contains(t)) continue;
    changes.push({ id: tag(t), step: stepOf(e), cls: e.dataset.addClass });
  }
  const videos = [...sl.querySelectorAll('video')].map((v) => {
    const r = v.getBoundingClientRect(), wrap = v.closest('[data-step]');
    return { src: v.getAttribute('src'), x: r.left, y: r.top, w: r.width, h: r.height,
             radius: parseFloat(getComputedStyle(v.closest('.video-wrap') || v).borderTopLeftRadius) || 0,
             chapters: v.dataset.chapters ? v.dataset.chapters.split(',').map(Number) : null,
             step: wrap && !wrap.classList.contains('step') ? stepOf(wrap) : (wrap ? stepOf(wrap) : 0) };
  });
  const chapterSteps = [...sl.querySelectorAll('[data-chapter]')].map((e) => ({ step: stepOf(e), chapter: +e.dataset.chapter }));
  const notes = [...sl.querySelectorAll('aside.notes p')].map((p) => ({ step: stepOf(p), text: p.textContent.replace(/\s+/g, ' ').trim() }));
  const steps = Math.max(0, ...[...sl.querySelectorAll('[data-step]')].map(stepOf));
  return { title: sl.dataset.title || '', steps, layers, changes, videos, chapterSteps, notes };
})()
"""

CAPTURE_CSS = """
html.cap *, html.cap *::before, html.cap *::after { transition: none !important; animation: none !important; }
html.cap #progress, html.cap #start, html.cap #blackout { display: none !important; }
html.cap video, html.cap .video-wrap .missing { visibility: hidden !important; }
html.cap .base-hide, html.cap .base-hide * { visibility: hidden !important; }
html.cap.layer, html.cap.layer body, html.cap.layer #stage, html.cap.layer .slide { background: transparent !important; }
html.cap.layer .slide.active *, html.cap.layer .slide.active::before, html.cap.layer .slide.active::after { visibility: hidden !important; }
html.cap.layer .slide.active .cap-on, html.cap.layer .slide.active .cap-on * { visibility: visible !important; }
html.cap.layer .slide.active .cap-on .cap-off, html.cap.layer .slide.active .cap-on .cap-off * { visibility: hidden !important; }
html.cap.layer .slide.active video { visibility: hidden !important; }
"""

LAYER_JS = r"""
([id, offIds]) => {
  const sl = document.querySelector('.slide.active');
  const el = sl.querySelector(`[data-cap-id="${id}"]`);
  document.documentElement.classList.add('layer');
  el.classList.add('cap-on');
  offIds.forEach((o) => sl.querySelector(`[data-cap-id="${o}"]`).classList.add('cap-off'));
  // nested steps of another click are not part of this picture
  const k = el.closest('[data-step]') ? +el.closest('[data-step]').dataset.step : 0;
  el.querySelectorAll('[data-step]').forEach((d) => { if (+d.dataset.step !== k) d.classList.add('cap-off'); });
  // the picture's box: the element and everything visible in it (doodles may overflow), plus room for shadows
  const boxes = [el, ...el.querySelectorAll('*')].filter((d) => !d.closest('.cap-off') && getComputedStyle(d).visibility === 'visible')
    .map((d) => d.getBoundingClientRect()).filter((r) => r.width > 0 && r.height > 0);
  const x0 = Math.min(...boxes.map((r) => r.left)), y0 = Math.min(...boxes.map((r) => r.top));
  const x1 = Math.max(...boxes.map((r) => r.right)), y1 = Math.max(...boxes.map((r) => r.bottom));
  const pad = 24;
  const x = Math.max(0, Math.floor(x0 - pad)), y = Math.max(0, Math.floor(y0 - pad));
  return { x, y, w: Math.min(1920, Math.ceil(x1 + pad)) - x, h: Math.min(1080, Math.ceil(y1 + pad)) - y };
}
"""

UNLAYER_JS = """() => { document.documentElement.classList.remove('layer');
  document.querySelectorAll('.cap-on, .cap-off').forEach((e) => e.classList.remove('cap-on', 'cap-off')); }"""


def clip(b):
    return {"x": b["x"], "y": b["y"], "width": b["w"], "height": b["h"]}


async def capture(tmp: Path):
    from playwright.async_api import async_playwright
    url = (PITCH / "index.html").as_uri()
    data = []
    async with async_playwright() as p:
        b = await p.chromium.launch(executable_path=CHROME if Path(CHROME).exists() else None)
        pg = await b.new_page(viewport={"width": 1920, "height": 1080})
        pg.on("pageerror", lambda e: print("pageerror:", e))
        await pg.goto(url + "?nostart")
        n_slides = await pg.evaluate("document.querySelectorAll('#stage .slide').length")
        for i in range(n_slides):
            await pg.goto(f"{url}?slide={i + 1}&step=max")
            await pg.evaluate("document.fonts.ready")
            await pg.add_style_tag(content=CAPTURE_CSS)
            await pg.evaluate("document.documentElement.classList.add('cap')")
            await pg.wait_for_timeout(900)   # images and the pen's last strokes
            info = await pg.evaluate(PAGE_JS)
            d = tmp / f"s{i + 1:02d}"; d.mkdir()
            # elements changed by a later click start in their first look
            await pg.evaluate("(info) => info.changes.forEach((c) => document.querySelector(`[data-cap-id='${c.id}']`).classList.remove(c.cls))", info)
            # layers: each animated element alone on a transparent page
            for L in info["layers"]:
                off = [o["id"] for o in info["layers"] if o is not L and o["step"] == L["step"] and o["anim"] == "draw" and o["id"] != L["id"]]
                L["box"] = await pg.evaluate(LAYER_JS, [L["id"], off])
                L["png"] = str(d / f"l{L['id']}.png")
                await pg.screenshot(path=L["png"], clip=clip(L["box"]), omit_background=True)
                await pg.evaluate(UNLAYER_JS)
            # an element whose look changes on a click: once per click, in its look at that click
            for C in info["changes"]:
                cls = [c["cls"] for c in info["changes"] if c["id"] == C["id"] and c["step"] <= C["step"]]
                allc = [c["cls"] for c in info["changes"] if c["id"] == C["id"]]
                await pg.evaluate("([id, rm, add]) => { const e = document.querySelector(`[data-cap-id='${id}']`); rm.forEach((c) => e.classList.remove(c)); add.forEach((c) => e.classList.add(c)); }", [C["id"], allc, cls])
                C["box"] = await pg.evaluate(LAYER_JS, [C["id"], []])
                C["png"] = str(d / f"c{C['id']}-{C['step']}.png")
                await pg.screenshot(path=C["png"], clip=clip(C["box"]), omit_background=True)
                await pg.evaluate(UNLAYER_JS)
            # background: the slide before any click, without its animated parts and videos
            await pg.evaluate("""(info) => {
              const sl = document.querySelector('.slide.active');
              info.layers.forEach((l) => sl.querySelector(`[data-cap-id='${l.id}']`).classList.add('base-hide'));
              info.changes.forEach((c) => sl.querySelector(`[data-cap-id='${c.id}']`).classList.remove(c.cls));
            }""", info)
            info["bg"] = str(d / "bg.png")
            await pg.screenshot(path=info["bg"])
            data.append(info)
            print(f"slide {i + 1:2d}  {info['title']}: {len(info['layers'])} layers, {info['steps']} clicks")
        await b.close()
    return data


# --------------------------------------------------------------------------- videos
def chapter_clips(src: Path, ends, tmp: Path):
    """One clip per chapter, as the deck plays them: chapter n runs from the end of n-1 to its own end."""
    at = lambda n: 0.0 if n <= 0 else ends[n - 1] - 0.08
    out = []
    for n in range(1, len(ends) + 1):
        clip, poster = tmp / f"chapter{n}.mp4", tmp / f"chapter{n}.png"
        subprocess.run(["ffmpeg", "-v", "error", "-y", "-ss", f"{at(n - 1):.3f}", "-to", f"{at(n):.3f}", "-i", str(src), "-an",
                        "-c:v", "libopenh264", "-b:v", "2M", "-pix_fmt", "yuv420p", "-movflags", "+faststart", str(clip)], check=True)
        subprocess.run(["ffmpeg", "-v", "error", "-y", "-i", str(clip), "-frames:v", "1", str(poster)], check=True)
        out.append((clip, poster, at(n) - at(n - 1)))
    return out


def poster_of(src: Path, tmp: Path):
    poster = tmp / (src.stem + "-poster.png")
    subprocess.run(["ffmpeg", "-v", "error", "-y", "-i", str(src), "-frames:v", "1", str(poster)], check=True)
    dur = float(subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", str(src)],
                               capture_output=True, text=True, check=True).stdout)
    return poster, dur


# --------------------------------------------------------------------------- PowerPoint animation XML
NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
PRESET = {  # kind: (presetID, presetSubtype, presetClass)
    "appear": (1, 0, "entr"), "fade": (10, 0, "entr"), "rise": (42, 0, "entr"), "pop": (53, 16, "entr"),
    "wipe": (22, 2, "entr"), "wipe-out": (22, 8, "exit"), "media": (1, 0, "mediacall"),
}


class Timing:
    def __init__(self):
        self.n = 2   # 1 = root, 2 = main sequence

    def id(self):
        self.n += 1
        return self.n

    def effect(self, spid, kind, delay, dur, node, grp=0):
        pid, sub, cls = PRESET[kind]
        i = self.id()
        tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
        head = f'<p:par><p:cTn id="{i}" presetID="{pid}" presetClass="{cls}" presetSubtype="{sub}" fill="hold" grpId="{grp}" nodeType="{node}"><p:stCondLst><p:cond delay="{int(delay)}"/></p:stCondLst><p:childTnLst>'
        tail = '</p:childTnLst></p:cTn></p:par>'
        if kind == "media":
            return head + f'<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="{self.id()}" dur="{int(dur)}" fill="hold"/>{tgt}</p:cBhvr></p:cmd>' + tail
        vis = lambda val, at=0: (f'<p:set><p:cBhvr><p:cTn id="{self.id()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="{at}"/></p:stCondLst></p:cTn>{tgt}'
                                 f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="{val}"/></p:to></p:set>')
        fx = lambda filt, d, how="in": f'<p:animEffect transition="{how}" filter="{filt}"><p:cBhvr><p:cTn id="{self.id()}" dur="{int(d)}"/>{tgt}</p:cBhvr></p:animEffect>'
        def move(attr, frm, d):
            return (f'<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn id="{self.id()}" dur="{int(d)}" decel="100000" fill="hold"/>{tgt}'
                    f'<p:attrNameLst><p:attrName>{attr}</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst>'
                    f'<p:tav tm="0"><p:val><p:strVal val="{frm}"/></p:val></p:tav><p:tav tm="100000"><p:val><p:strVal val="#{attr}"/></p:val></p:tav></p:tavLst></p:anim>')
        if kind == "appear":
            body = vis("visible")
        elif kind == "fade":
            body = vis("visible") + fx("fade", dur)
        elif kind == "rise":
            body = vis("visible") + fx("fade", dur) + move("ppt_y", "#ppt_y+0.024", dur * 1.25)
        elif kind == "pop":
            body = vis("visible") + fx("fade", dur * 0.5) + move("ppt_w", "#ppt_w*0.6", dur) + move("ppt_h", "#ppt_h*0.6", dur)
        elif kind == "wipe":   # right to left, the way Arabic is written and the pen draws
            body = vis("visible") + fx("wipe(right)", dur)
        elif kind == "wipe-out":   # erased from the end of the line (the left) back to its start
            body = fx("wipe(left)", dur, "out") + vis("hidden", int(dur) - 1)
        return head + body + tail

    def group(self, effects_fn, auto=False):
        outer, inner = self.id(), self.id()
        cond = '<p:cond delay="indefinite"/>' + ('<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond>' if auto else "")
        body = effects_fn()
        return (f'<p:par><p:cTn id="{outer}" fill="hold"><p:stCondLst>{cond}</p:stCondLst><p:childTnLst>'
                f'<p:par><p:cTn id="{inner}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{body}'
                f'</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')

    def xml(self, groups, videos):
        seq = ""
        if groups:
            seq = (f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{"".join(groups)}</p:childTnLst></p:cTn>'
                   '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
                   '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>')
        media = "".join(f'<p:video><p:cMediaNode vol="0" mute="1"><p:cTn id="{self.id()}" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn>'
                        f'<p:tgtEl><p:spTgt spid="{v}"/></p:tgtEl></p:cMediaNode></p:video>' for v in videos)
        return f'<p:timing {NS}><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>{seq}{media}</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'


DUR = {"rise": 600, "fade": 550, "pop": 600, "draw": 1000, "appear": 1}


# --------------------------------------------------------------------------- build
def build(data, tmp: Path):
    from lxml import etree
    from pptx import Presentation
    from pptx.util import Emu
    from pptx.oxml.ns import qn
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(1920 * PX), Emu(1080 * PX)
    blank = prs.slide_layouts[6]
    px = lambda v: Emu(int(round(v * PX)))

    for info in data:
        s = prs.slides.add_slide(blank)
        bg = s.shapes.add_picture(info["bg"], 0, 0, prs.slide_width, prs.slide_height)
        bg._element.nvPicPr.cNvPr.set("descr", info["title"])
        bg.name = "Slide"
        pics = {}

        def pic(path, box, name):
            sh = s.shapes.add_picture(path, px(box["x"]), px(box["y"]), px(box["w"]), px(box["h"]))
            sh.name = name
            return sh.shape_id

        clicks = {k: [] for k in range(info["steps"] + 1)}   # k -> [(spid, kind, delay, dur)]
        auto = []
        # videos sit under the step layers
        vids = []
        for v in info["videos"]:
            src = PITCH / v["src"]
            box = {"x": v["x"], "y": v["y"], "w": v["w"], "h": v["h"]}
            if v["chapters"]:
                clips = chapter_clips(src, v["chapters"], tmp)
                chap = {c["chapter"]: c["step"] for c in info["chapterSteps"]}
                for n, (clip, poster, dur) in enumerate(clips, 1):
                    m = s.shapes.add_movie(str(clip), px(box["x"]), px(box["y"]), px(box["w"]), px(box["h"]), poster_frame_image=str(poster), mime_type="video/mp4")
                    m.name = f"Demo chapter {n}"
                    vids.append((m, v["radius"]))
                    k = chap.get(n, n)
                    if n > 1:
                        clicks[k].append((m.shape_id, "appear", 0, 1))
                    clicks[k].append((m.shape_id, "media", 0, int(dur * 1000) + 200))
            else:
                poster, dur = poster_of(src, tmp)
                m = s.shapes.add_movie(str(src), px(box["x"]), px(box["y"]), px(box["w"]), px(box["h"]), poster_frame_image=str(poster), mime_type="video/mp4")
                m.name = "Demo, uncut"
                vids.append((m, v["radius"]))
                clicks[max(1, v["step"])].append((m.shape_id, "media", 0, int(dur * 1000) + 200))

        layers = sorted(info["layers"], key=lambda L: (not L.get("auto"), L["step"]))
        for L in layers:
            spid = pic(L["png"], L["box"], f"Step {L['step']}")
            delay = L["delay"] * 1000
            target = auto if L.get("auto") else clicks[L["step"]]
            if L["anim"] in ("type", "type-erase"):
                d_in = max(600, L["chars"] * 45)
                target.append((spid, "wipe", delay, d_in))
                if L["anim"] == "type-erase":
                    target.append((spid, "wipe-out", delay + d_in + 1500, max(500, L["chars"] * 38), 1))
            elif L["anim"] == "draw":
                target.append((spid, "wipe", delay + (250 if L.get("auto") else 0), DUR["draw"]))
            else:
                kind = L["anim"] if L["anim"] in ("fade", "pop") else "rise"
                target.append((spid, kind, delay, DUR[kind]))
        for C in info["changes"]:
            clicks[C["step"]].append((pic(C["png"], C["box"], f"Step {C['step']} change"), "fade", 0, 500))

        # clicks that only move the script still take a click, as in the HTML deck
        for k in range(1, info["steps"] + 1):
            if not clicks[k]:
                dot = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, px(1), px(1))
                dot.fill.background(); dot.line.fill.background(); dot.name = f"Step {k} (script only)"
                clicks[k].append((dot.shape_id, "appear", 0, 1))

        # rounded corners on the videos, like the phone frame around them
        for m, r in vids:
            geom = m._element.spPr.find(qn("a:prstGeom"))
            if geom is not None and r:
                geom.set("prst", "roundRect")
                av = geom.find(qn("a:avLst"))
                if av is None:
                    av = etree.SubElement(geom, qn("a:avLst"))
                gd = etree.SubElement(av, qn("a:gd")); gd.set("name", "adj"); gd.set("fmla", f"val {int(r / min(m.width, m.height) * PX * 100000)}")

        t = Timing()
        def effects(items, first_node):
            def f():
                return "".join(t.effect(it[0], it[1], it[2], it[3], first_node if j == 0 else "withEffect", it[4] if len(it) > 4 else 0)
                               for j, it in enumerate(items))
            return f
        groups = []
        if auto:
            groups.append(t.group(effects(auto, "afterEffect"), auto=True))
        for k in range(1, info["steps"] + 1):
            groups.append(t.group(effects(clicks[k], "clickEffect")))
        sld = s._element
        old = sld.find(qn("p:timing"))
        if old is not None:
            sld.remove(old)
        trans = etree.fromstring(f'<p:transition {NS} spd="med"><p:fade/></p:transition>')
        timing = etree.fromstring(t.xml(groups, [m.shape_id for m, _ in vids]))
        ext = sld.find(qn("p:extLst"))
        for el in (trans, timing):
            if ext is not None:
                ext.addprevious(el)
            else:
                sld.append(el)

        # the script, one paragraph per line, with the click it belongs to
        tf = s.notes_slide.notes_text_frame
        tf.text = ""
        last = None
        for j, line in enumerate(info["notes"]):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            mark = f"[{line['step']}] " if line["step"] != last else ""
            last = line["step"]
            p.text = mark + line["text"]
            p._p.get_or_add_pPr().set("rtl", "1")
            p._p.get_or_add_pPr().set("algn", "r")

    prs.core_properties.title = "خطوة · Khutwa"
    prs.core_properties.author = "Khutwa team"
    prs.save(OUT)
    print(f"{OUT.relative_to(PITCH.parent)}: {OUT.stat().st_size / 1e6:.1f} MB, {len(data)} slides")


def main():
    tmp = Path(tempfile.mkdtemp(prefix="khutwa-pptx-"))
    try:
        data = asyncio.run(capture(tmp))
        build(data, tmp)
    finally:
        if "--keep" in sys.argv:
            print("layers kept in", tmp)
        else:
            shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    main()
